#!/usr/bin/env python3
"""
Generate master-product.js dari Master Product.xlsx.

Pakai:
    python tools/build_master_product.py "Master Product.xlsx"
    python tools/build_master_product.py "Master Product.xlsx" --out master-product.js

Kolom yang dibaca (header baris 1, nama tidak case-sensitive):
    Category | Subcategory | Product Code | Product Name | Unit

Output: satu file JS (single source of truth) yang di-load website secara
asynchronous. Format kompak: array [category, subcategory, code, name, unit].
"""
import argparse
import datetime as dt
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Butuh openpyxl: pip install openpyxl")

REQUIRED = ["category", "subcategory", "product code", "product name", "unit"]


def clean(v):
    if v is None:
        return ""
    return " ".join(str(v).split()).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--out", default=str(Path(__file__).resolve().parent.parent / "master-product.js"))
    ap.add_argument("--sheet", default=None, help="Nama sheet (default: sheet pertama)")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, read_only=True, data_only=True)
    ws = wb[a.sheet] if a.sheet else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [clean(h).lower() for h in next(rows)]
    idx = {}
    for need in REQUIRED:
        if need not in header:
            sys.exit(f"Kolom '{need}' tidak ditemukan. Header: {header}")
        idx[need] = header.index(need)

    items, seen_codes, dup_names, skipped = [], set(), {}, 0
    for r in rows:
        if r is None or all(clean(c) == "" for c in r):
            continue
        cat, sub, code, name, unit = (clean(r[idx[k]]) for k in REQUIRED)
        if not name:
            skipped += 1
            continue
        if code and code in seen_codes:
            skipped += 1
            continue
        if code:
            seen_codes.add(code)
        dup_names[name.upper()] = dup_names.get(name.upper(), 0) + 1
        items.append([cat.upper(), sub.upper(), code, name, unit])

    items.sort(key=lambda x: (x[3].upper(), x[2]))
    dups = sum(1 for n in dup_names.values() if n > 1)
    cats = sorted({i[0] for i in items})
    subs = sorted({i[1] for i in items})

    body = ",\n".join(json.dumps(i, ensure_ascii=False) for i in items)
    js = (
        "/* =========================================================\n"
        "   MASTER PRODUCT — di-generate otomatis, JANGAN edit manual.\n"
        f"   Sumber : {Path(a.xlsx).name}\n"
        f"   Dibuat : {dt.datetime.now():%Y-%m-%d %H:%M}\n"
        f"   Produk : {len(items)} | Kategori: {len(cats)} | Subkategori: {len(subs)}\n"
        "   Update : python tools/build_master_product.py \"Master Product.xlsx\"\n"
        "   Format : [Category, Subcategory, Product Code, Product Name, Unit]\n"
        "   ========================================================= */\n"
        "window.MASTER_PRODUCT_DATA = {\n"
        f"    generatedAt: {json.dumps(dt.datetime.now().strftime('%Y-%m-%d %H:%M'))},\n"
        f"    source: {json.dumps(Path(a.xlsx).name)},\n"
        "    columns: [\"category\", \"subcategory\", \"code\", \"name\", \"unit\"],\n"
        "    rows: [\n" + body + "\n    ]\n};\n"
        "if (typeof MasterProduct !== 'undefined' && MasterProduct._receive) MasterProduct._receive(window.MASTER_PRODUCT_DATA);\n"
    )
    Path(a.out).write_text(js, encoding="utf-8")
    print(f"OK  {a.out}")
    print(f"    {len(items)} produk, {len(cats)} kategori, {len(subs)} subkategori")
    print(f"    nama duplikat: {dups} (ditampilkan dengan kode & unit sebagai pembeda), baris dilewati: {skipped}")


if __name__ == "__main__":
    main()
